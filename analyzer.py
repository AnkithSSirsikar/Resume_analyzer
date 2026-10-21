"""Core logic for the Resume Analyzer: text extraction + scoring."""
import io
import re
from collections import Counter
from dataclasses import dataclass

from sklearn.feature_extraction.text import ENGLISH_STOP_WORDS, TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# --------------------------------------------------------------------------
# 1. Skill dictionary: canonical name -> accepted spellings.
#    Add more skills here to make the analyzer smarter for your domain.
# --------------------------------------------------------------------------
SKILLS = {
    "python": ["python"],
    "java": ["java"],
    "javascript": ["javascript", "js", "ecmascript"],
    "typescript": ["typescript"],
    "c++": ["c++"],
    "c#": ["c#"],
    "go": ["golang"],
    "rust": ["rust"],
    "kotlin": ["kotlin"],
    "swift": ["swift"],
    "php": ["php"],
    "sql": ["sql"],
    "mysql": ["mysql"],
    "postgresql": ["postgresql", "postgres"],
    "mongodb": ["mongodb", "mongo"],
    "redis": ["redis"],
    "firebase": ["firebase"],
    "html": ["html", "html5"],
    "css": ["css", "css3"],
    "react": ["react", "reactjs", "react.js"],
    "angular": ["angular", "angularjs"],
    "vue": ["vue", "vuejs", "vue.js"],
    "node.js": ["node.js", "nodejs", "node"],
    "express": ["express", "express.js"],
    "django": ["django"],
    "flask": ["flask"],
    "fastapi": ["fastapi"],
    "spring boot": ["spring boot", "springboot", "spring"],
    "rest api": ["rest api", "rest apis", "restful", "rest"],
    "graphql": ["graphql"],
    "microservices": ["microservices", "microservice"],
    "machine learning": ["machine learning", "ml"],
    "deep learning": ["deep learning"],
    "nlp": ["nlp", "natural language processing"],
    "computer vision": ["computer vision", "opencv"],
    "tensorflow": ["tensorflow"],
    "pytorch": ["pytorch"],
    "scikit-learn": ["scikit-learn", "sklearn"],
    "pandas": ["pandas"],
    "numpy": ["numpy"],
    "data analysis": ["data analysis", "data analytics"],
    "data structures": ["data structures"],
    "algorithms": ["algorithms", "dsa"],
    "oop": ["oop", "object oriented", "object-oriented"],
    "git": ["git"],
    "github": ["github"],
    "docker": ["docker"],
    "kubernetes": ["kubernetes", "k8s"],
    "aws": ["aws", "amazon web services"],
    "azure": ["azure"],
    "gcp": ["gcp", "google cloud"],
    "linux": ["linux"],
    "ci/cd": ["ci/cd", "cicd"],
    "android": ["android"],
    "flutter": ["flutter"],
    "tableau": ["tableau"],
    "power bi": ["power bi", "powerbi"],
    "excel": ["excel"],
    "jira": ["jira"],
    "agile": ["agile", "scrum"],
    "testing": ["unit testing", "pytest", "junit", "selenium"],
    "communication": ["communication"],
    "teamwork": ["teamwork", "team player", "collaboration"],
    "leadership": ["leadership"],
    "problem solving": ["problem solving", "problem-solving"],
}


def _compile(variants):
    parts = "|".join(re.escape(v) for v in variants)
    return re.compile(r"(?<![a-z0-9+#])(?:" + parts + r")(?![a-z0-9+#])")


SKILL_PATTERNS = {name: _compile(variants) for name, variants in SKILLS.items()}


def find_skills(text):
    t = text.lower()
    return {name for name, pat in SKILL_PATTERNS.items() if pat.search(t)}


# --------------------------------------------------------------------------
# 2. Generic keyword extraction from the job description
# --------------------------------------------------------------------------
GENERIC_WORDS = {
    "experience", "work", "working", "ability", "team", "strong", "skills", "skill",
    "looking", "role", "job", "candidate", "candidates", "required", "requirements",
    "responsibilities", "years", "year", "will", "must", "etc", "good", "great",
    "excellent", "knowledge", "understanding", "preferred", "plus", "including",
    "company", "position", "join", "opportunity", "using", "use", "well", "new",
    "like", "least", "related", "relevant", "environment", "based", "across",
    "ensure", "help", "highly", "able", "need", "needs", "want", "seeking",
}
STOP = set(ENGLISH_STOP_WORDS) | GENERIC_WORDS


def top_keywords(jd_text, n=15):
    tokens = re.findall(r"[a-z][a-z+#]*", jd_text.lower())
    valid = lambda w: len(w) > 2 and w not in STOP
    counts = Counter(w for w in tokens if valid(w))
    for a, b in zip(tokens, tokens[1:]):
        if valid(a) and valid(b):
            counts[f"{a} {b}"] += 1
    # drop bigrams that appear only once (usually noise) and anything already a known skill
    ranked = [
        (k, c) for k, c in counts.most_common()
        if (" " not in k or c >= 2) and not find_skills(k)
    ]
    return [k for k, _ in ranked[:n]]


def _in_text(keyword, text_lower):
    # prefix match on a lightly-stemmed word so "databases" matches "database"
    stem = keyword.rstrip("s")
    return re.search(r"\b" + re.escape(stem), text_lower) is not None


# --------------------------------------------------------------------------
# 3. Text similarity (TF-IDF + cosine)
# --------------------------------------------------------------------------
def similarity(resume_text, jd_text):
    try:
        vec = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), sublinear_tf=True)
        m = vec.fit_transform([resume_text, jd_text])
        return float(cosine_similarity(m[0], m[1])[0][0])
    except ValueError:  # empty vocabulary
        return 0.0


# --------------------------------------------------------------------------
# 4. Resume quality checks (independent of the job description)
# --------------------------------------------------------------------------
SECTION_PATTERNS = {
    "education": r"^\s*(education|academic)",
    "skills": r"^\s*(technical\s+)?skills|^\s*core competencies",
    "projects": r"^\s*(academic\s+|personal\s+)?projects?",
    "experience": r"^\s*(work\s+|professional\s+)?experience|^\s*internships?|^\s*employment",
    "summary": r"^\s*(professional\s+)?(summary|objective|profile)|^\s*about me",
}
SECTION_REGEX = {k: re.compile(v, re.I | re.M) for k, v in SECTION_PATTERNS.items()}

ACTION_VERBS = {
    "developed", "built", "designed", "implemented", "created", "led", "managed",
    "optimized", "improved", "reduced", "increased", "automated", "deployed",
    "analyzed", "engineered", "launched", "integrated", "collaborated", "delivered",
    "architected", "trained", "tested", "debugged", "migrated", "streamlined",
    "mentored", "achieved", "researched", "organized", "coordinated", "established",
}
WEAK_PHRASES = ["responsible for", "worked on", "helped with", "duties included", "tasked with"]

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(\+?\d[\d\s\-()]{8,}\d)")
LINK_RE = re.compile(r"linkedin\.com|github\.com", re.I)
METRIC_RE = re.compile(
    r"\b\d+(?:\.\d+)?\s?%"
    r"|[$₹]\s?\d"
    r"|\b\d[\d,]*\+?\s+(?:users|students|records|requests|customers|clients|projects|members|downloads)",
    re.I,
)


def format_check(resume_text):
    """Returns (score 0-100, sections_found, suggestions, stats)."""
    suggestions = []
    score = 0.0
    words = re.findall(r"\b\w+\b", resume_text.lower())
    word_count = len(words)

    # contact details (20 pts)
    has_email = bool(EMAIL_RE.search(resume_text))
    has_phone = bool(PHONE_RE.search(resume_text))
    has_link = bool(LINK_RE.search(resume_text))
    score += 10 * has_email + 5 * has_phone + 5 * has_link
    if not has_email:
        suggestions.append("Add a professional email address at the top of your resume.")
    if not has_phone:
        suggestions.append("Add a phone number so recruiters can reach you.")
    if not has_link:
        suggestions.append("Add your LinkedIn and/or GitHub link - recruiters for tech roles check them.")

    # sections (40 pts)
    sections = {name: bool(rx.search(resume_text)) for name, rx in SECTION_REGEX.items()}
    score += 10 * sections["education"] + 10 * sections["skills"] + 10 * sections["projects"]
    if sections["experience"]:
        score += 10
    elif sections["projects"]:
        score += 5  # freshers get partial credit if they show projects instead
    for name in ("education", "skills", "projects"):
        if not sections[name]:
            suggestions.append(f"Add a clear '{name.title()}' section heading.")
    if not sections["experience"]:
        suggestions.append(
            "Add an 'Experience' or 'Internships' section (even short internships, "
            "freelance work or open-source contributions count)."
        )
    if not sections["summary"]:
        suggestions.append("Consider a 2-3 line 'Summary' tailored to this specific job.")

    # length (10 pts)
    if 300 <= word_count <= 900:
        score += 10
    elif 200 <= word_count < 300 or 900 < word_count <= 1200:
        score += 5
    if word_count < 300:
        suggestions.append(f"Your resume is short ({word_count} words). Add detail about projects and impact.")
    elif word_count > 900:
        suggestions.append(f"Your resume is long ({word_count} words). Aim for 1 page (2 at most).")

    # action verbs (15 pts)
    verb_count = sum(1 for w in words if w in ACTION_VERBS)
    score += min(verb_count / 8, 1) * 15
    if verb_count < 8:
        suggestions.append(
            "Start bullet points with strong action verbs (built, designed, optimized, deployed...)."
        )

    # quantified achievements (15 pts)
    metric_count = len(METRIC_RE.findall(resume_text))
    score += min(metric_count / 3, 1) * 15
    if metric_count < 3:
        suggestions.append(
            "Quantify your impact, e.g. 'reduced load time by 35%' or 'served 500+ users'."
        )

    # weak phrases
    low = resume_text.lower()
    weak_found = [p for p in WEAK_PHRASES if p in low]
    if weak_found:
        suggestions.append(
            "Replace weak phrases (" + ", ".join(f"'{p}'" for p in weak_found)
            + ") with specific achievements."
        )

    stats = {
        "word_count": word_count,
        "action_verbs": verb_count,
        "quantified_achievements": metric_count,
    }
    return round(score), sections, suggestions, stats


# --------------------------------------------------------------------------
# 5. Main entry point
# --------------------------------------------------------------------------
@dataclass
class AnalysisResult:
    overall: int
    breakdown: dict
    matched_skills: list
    missing_skills: list
    missing_keywords: list
    sections_found: dict
    suggestions: list
    stats: dict


def analyze(resume_text, jd_text):
    resume_text = resume_text.strip()
    jd_text = jd_text.strip()
    rl = resume_text.lower()

    jd_skills = find_skills(jd_text)
    resume_skills = find_skills(resume_text)
    matched = sorted(jd_skills & resume_skills)
    missing = sorted(jd_skills - resume_skills)

    keywords = top_keywords(jd_text)
    hits = [k for k in keywords if _in_text(k, rl)]
    missing_kw = [k for k in keywords if k not in hits]
    keyword_score = len(hits) / len(keywords) * 100 if keywords else 0.0

    skill_score = len(matched) / len(jd_skills) * 100 if jd_skills else keyword_score

    sim = similarity(resume_text, jd_text)
    sim_score = min(sim / 0.45, 1.0) * 100  # ~0.45 cosine is already a very strong match

    fmt_score, sections, fmt_suggestions, stats = format_check(resume_text)

    overall = 0.40 * skill_score + 0.20 * keyword_score + 0.15 * sim_score + 0.25 * fmt_score

    suggestions = []
    if missing:
        suggestions.append(
            "The job asks for skills not found in your resume: " + ", ".join(missing)
            + ". Add the ones you genuinely have (in Skills and in project bullets)."
        )
    if missing_kw:
        suggestions.append(
            "Frequent job-description terms missing from your resume: "
            + ", ".join(missing_kw[:10]) + ". Use them naturally where they are true."
        )
    if sim_score < 50:
        suggestions.append(
            "Overall wording is far from the job description. Mirror its language in your "
            "summary and project descriptions."
        )
    suggestions.extend(fmt_suggestions)

    return AnalysisResult(
        overall=round(overall),
        breakdown={
            "Skill match": round(skill_score),
            "Keyword coverage": round(keyword_score),
            "Content similarity": round(sim_score),
            "Resume quality": round(fmt_score),
        },
        matched_skills=matched,
        missing_skills=missing,
        missing_keywords=missing_kw,
        sections_found=sections,
        suggestions=suggestions,
        stats=stats,
    )


# --------------------------------------------------------------------------
# 6. Reading uploaded files
# --------------------------------------------------------------------------
def extract_text(data: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if name.endswith(".docx"):
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    return data.decode("utf-8", errors="ignore")
